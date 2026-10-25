"""Sneak attempts A12-A15 and a harmless CNC program: drawings hidden in everyday file types.

  A12  Word RFQ with the label-cropped drawing pasted in           -> fingerprint, to Germany
  A13  PowerPoint quote deck with a screenshot of the drawing       -> fingerprint, to Germany
  A14  ZIP job folder, the drawing two levels down in an inner ZIP  -> fingerprint, to Germany
  A15  CNC program whose header names an ITAR part                  -> export register, to Germany
  KP-1047_op10.nc  CNC program for a commercial part                -> released to Triton Marine

Dev-only dependencies (the engine itself only needs zipfile): uv pip install python-docx python-pptx
Appends rows to data/outgoing/answer_key.csv once; safe to re-run.
"""
from __future__ import annotations

import csv
import io
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from exportgate.config import SETTINGS  # noqa: E402

D = SETTINGS.data_dir
ATT = D / "attacks"
ENG = "marco.bellini@kestrelprecision.example"
DE = "hans.weber@mueller-praezision.example"
TRITON = "sales@tritonmarine.example"


def gcode(part: str, desc: str, matl: str) -> str:
    moves = "\n".join(f"G01 X{10 + i * 4.5:.3f} Y{(i % 5) * 3.2:.3f} F{600 + i * 10}" for i in range(24))
    return f"""%
O4471 ({part} {desc} OP10)
(KESTREL PRECISION - HAAS VF-2)
(MATL {matl})
(T1 1/2 FLAT END MILL  T2 #7 DRILL)
G20 G90 G17 G40 G80
T1 M06
G54 G00 X0. Y0. S8000 M03
G43 H01 Z1.0 M08
{moves}
G00 Z1.0 M09
M05
G28 G91 Z0.
M30
%
"""


def main():
    from docx import Document
    from docx.shared import Inches
    from pptx import Presentation
    from pptx.util import Inches as PI, Pt

    crop = ATT / "A01_bracket_views_only.png"
    shot = ATT / "A03_screenshot.png"
    pdf = ATT / "A02_whiteout_markings.pdf"

    doc = Document()
    doc.add_heading("RFQ 4471 - machining quote", 1)
    doc.add_paragraph("Hi Hans,\n\nPlease quote 40 pcs of the bracket below, material 6061-T6, anodize type II. "
                      "Delivery 6 weeks. Views attached, let me know if you need the model.\n\nThanks, Marco")
    doc.add_picture(str(crop), width=Inches(6))
    doc.save(ATT / "A12_rfq_word.docx")

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Q4 supplier quote: bracket family"
    s.shapes.add_picture(str(shot), PI(0.6), PI(1.6), width=PI(8.8))
    tb = s.shapes.add_textbox(PI(0.6), PI(6.6), PI(8.8), PI(0.6)).text_frame
    tb.text = "Target price $38/pc at 500/yr. Tolerances per attached views."
    tb.paragraphs[0].runs[0].font.size = Pt(14)
    prs.save(ATT / "A13_quote_deck.pptx")

    inner = io.BytesIO()
    with zipfile.ZipFile(inner, "w", zipfile.ZIP_DEFLATED) as z:
        z.write(pdf, "old_revs/bracket_rev.pdf")
    with zipfile.ZipFile(ATT / "A14_job_folder.zip", "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("job_4471/notes.txt", "Job 4471. Customer wants first article by the 14th.\n")
        z.writestr("job_4471/archive.zip", inner.getvalue())
        z.write(crop, "job_4471/views.png")

    (ATT / "A15_op10_program.nc").write_text(gcode("KP-1023", "SENSOR BRACKET", "6061-T6"))
    (D / "outgoing" / "KP-1047_op10.nc").write_text(gcode("KP-1047", "SERVO HOUSING", "6061-T6"))

    key = D / "outgoing" / "answer_key.csv"
    rows = list(csv.DictReader(open(key)))
    have = {r["file"] for r in rows}
    new = [
        ("attack", "attacks/A12_rfq_word.docx", "office_docx", "KP-1023", "ITAR", 1, ENG, DE, "cropped drawing pasted into a Word RFQ"),
        ("attack", "attacks/A13_quote_deck.pptx", "office_pptx", "KP-1023", "ITAR", 1, ENG, DE, "screenshot of the drawing in a quote deck"),
        ("attack", "attacks/A14_job_folder.zip", "zip_nested", "KP-1023", "ITAR", 1, ENG, DE, "drawing inside a ZIP, plus a ZIP inside the ZIP"),
        ("attack", "attacks/A15_op10_program.nc", "cnc_program", "KP-1023", "ITAR", 1, ENG, DE, "CNC program; the header names the ITAR part"),
        ("outgoing", "outgoing/KP-1047_op10.nc", "cnc_program", "KP-1047", "EAR99", 0, ENG, TRITON, "CNC program for a commercial part"),
    ]
    added = 0
    with open(key, "a", newline="") as f:
        w = csv.writer(f)
        for r in new:
            if r[1] not in have:
                w.writerow(r)
                added += 1
    print(f"made A12-A15 + KP-1047_op10.nc; {added} answer-key rows added")


if __name__ == "__main__":
    main()
